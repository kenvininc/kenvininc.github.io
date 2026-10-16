import io
import shutil
import tempfile
import uuid

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook

from apps.catalog.excel import HEADERS
from apps.catalog.models import PriceHistory, Product, ProductImage, StockMovement, Variant
from apps.core.models import AuditLog, SiteSettings
from apps.orders.models import Order
from apps.orders.services import place_order

from .helpers import PASSWORD, configure_site, customer, make_product, make_user, png_upload

MEDIA = tempfile.mkdtemp()


def formset_data(rows, prefix="v", initial=0):
    data = {f"{prefix}-TOTAL_FORMS": str(len(rows)), f"{prefix}-INITIAL_FORMS": str(initial),
            f"{prefix}-MIN_NUM_FORMS": "1", f"{prefix}-MAX_NUM_FORMS": "1000"}
    for i, r in enumerate(rows):
        for k, v in r.items():
            data[f"{prefix}-{i}-{k}"] = v
    return data


@override_settings(MEDIA_ROOT=MEDIA)
class PermissionTests(TestCase):
    def setUp(self):
        configure_site()
        self.p, (self.v,) = make_product(stock=5)
        self.order, _ = place_order({self.v.pk: 1}, customer(), "cod", "", uuid.uuid4())

    def login(self, *roles):
        u = make_user("u" + uuid.uuid4().hex[:6], *roles)
        self.client.login(username=u.username, password=PASSWORD)
        return u

    def test_anonymous_redirected(self):
        for url in ("/admin/", "/admin/san-pham/", "/admin/don-hang/", "/admin/cai-dat/"):
            r = self.client.get(url)
            self.assertEqual(r.status_code, 302, url)
            self.assertIn("/admin/dang-nhap/", r["Location"])
        r = self.client.post(f"/admin/don-hang/{self.order.pk}/trang-thai/", {"status": "confirmed"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Order.objects.get(pk=self.order.pk).status, "new")

    def test_order_role_cannot_touch_products_prices_settings_users(self):
        self.login("Xử lý đơn hàng")
        self.assertEqual(self.client.get("/admin/don-hang/").status_code, 200)
        forbidden = [
            ("get", "/admin/san-pham/them/"), ("post", "/admin/san-pham/thao-tac/"), ("post", "/admin/gia/hang-loat/"),
            ("post", f"/admin/ton-kho/{self.v.pk}/"), ("post", "/admin/cai-dat/"), ("post", "/admin/tai-khoan/them/"),
            ("post", "/admin/khuyen-mai/them/"), ("get", "/admin/excel/xuat/"), ("post", f"/admin/san-pham/{self.p.pk}/xoa/"),
        ]
        for method, url in forbidden:
            r = getattr(self.client, method)(url, {"action": "publish", "ids": [self.p.pk], "quantity": 1})
            self.assertEqual(r.status_code, 403, url)
        self.assertTrue(Product.objects.filter(pk=self.p.pk).exists())

    def test_content_role_cannot_process_orders(self):
        self.login("Nội dung")
        r = self.client.post(f"/admin/don-hang/{self.order.pk}/trang-thai/", {"status": "cancelled", "note": "x"})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.client.get("/admin/don-hang/").status_code, 403)
        r = self.client.post(f"/admin/don-hang/{self.order.pk}/thanh-toan/", {"payment_status": "paid"})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(Order.objects.get(pk=self.order.pk).status, "new")
        # Nội dung được đánh dấu nổi bật nhưng không được đổi trạng thái bán
        r = self.client.post("/admin/san-pham/thao-tac/", {"action": "hide", "ids": [self.p.pk]})
        self.assertEqual(Product.objects.get(pk=self.p.pk).status, "published")

    def test_product_role_cannot_see_orders_or_revenue(self):
        self.login("Quản lý sản phẩm")
        self.assertEqual(self.client.get("/admin/don-hang/").status_code, 403)
        r = self.client.get("/admin/")
        self.assertEqual(r.status_code, 200)
        self.assertNotContains(r, "Doanh thu trong kỳ")

    def test_non_staff_and_inactive_cannot_login(self):
        u = make_user("khoa", "Chủ shop", active=False)
        r = self.client.post("/admin/dang-nhap/", {"username": "khoa", "password": PASSWORD})
        self.assertContains(r, "không đúng")
        u.is_active, u.is_staff = True, False
        u.save()
        r = self.client.post("/admin/dang-nhap/", {"username": "khoa", "password": PASSWORD})
        self.assertContains(r, "không đúng")

    def test_login_rate_limit_and_no_password_in_log(self):
        make_user("chu", "Chủ shop")
        for _ in range(5):
            self.client.post("/admin/dang-nhap/", {"username": "chu", "password": "sai-mat-khau-123"})
        r = self.client.post("/admin/dang-nhap/", {"username": "chu", "password": PASSWORD})
        self.assertContains(r, "quá nhiều lần")
        self.assertNotIn("_auth_user_id", self.client.session)
        for log in AuditLog.objects.all():
            self.assertNotIn("sai-mat-khau", log.summary)
            self.assertNotIn(PASSWORD, log.summary)

    def test_owner_cannot_remove_last_owner(self):
        u = self.login("Chủ shop")
        r = self.client.post(f"/admin/tai-khoan/{u.pk}/", {"first_name": "", "email": "", "roles": ["Nội dung"], "is_active": "on"})
        self.assertContains(r, "không thể tự khóa")
        self.assertTrue(u.groups.filter(name="Chủ shop").exists())


@override_settings(MEDIA_ROOT=MEDIA)
class AdminProductFlowTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        configure_site()
        self.owner = make_user("owner", "Chủ shop")
        self.client.login(username="owner", password=PASSWORD)

    def create_product(self):
        data = {"name": "Son Kem Lì Thử", "status": "published", "short_description": "", **formset_data([
            {"value": "01 Đỏ", "sku": "SKT-01", "price": "159000", "compare_at_price": "199000", "stock": "4", "is_active": "on", "sort_order": "0"},
            {"value": "02 Cam", "sku": "SKT-02", "price": "169000", "compare_at_price": "", "stock": "0", "is_active": "on", "sort_order": "1"},
        ])}
        data["new_images"] = [png_upload("a.png"), png_upload("b.png")]
        r = self.client.post("/admin/san-pham/them/", data)
        self.assertEqual(r.status_code, 302, r.content.decode()[:3000] if r.status_code != 302 else "")
        return Product.objects.get(name="Son Kem Lì Thử")

    def test_create_publish_buy_flow(self):
        p = self.create_product()
        self.assertEqual(p.variants.count(), 2)
        self.assertEqual(Variant.objects.get(sku="SKT-01").stock, 4)
        self.assertEqual(p.images.count(), 2)
        self.assertEqual(p.images.filter(is_main=True).count(), 1)
        img = p.images.first()
        self.assertTrue(img.image.name.endswith(".webp"))
        self.assertEqual(StockMovement.objects.filter(variant__sku="SKT-01", kind="initial").count(), 1)
        self.assertEqual(PriceHistory.objects.filter(variant__product=p, source="create").count(), 2)

        # Khách tìm kiếm không dấu, xem sản phẩm, chọn biến thể, thêm giỏ, đặt hàng
        shop = self.client_class()
        r = shop.get("/tim-kiem/", {"q": "son kem li"})
        self.assertContains(r, "Son Kem Lì Thử")
        r = shop.get(p.get_absolute_url())
        self.assertContains(r, "159.000 ₫")
        self.assertContains(r, "199.000 ₫")
        v = Variant.objects.get(sku="SKT-01")
        r = shop.post("/gio-hang/them/", {"variant_id": v.pk, "quantity": 5}, HTTP_X_REQUESTED_WITH="fetch")
        self.assertEqual(r.status_code, 409)
        self.assertIn("Chỉ còn 4", r.json()["message"])
        r = shop.post("/gio-hang/them/", {"variant_id": v.pk, "quantity": 2}, HTTP_X_REQUESTED_WITH="fetch")
        self.assertTrue(r.json()["ok"])
        self.assertEqual(r.json()["cart_count"], 2)
        r = shop.get("/thanh-toan/")
        key = r.context["form"]["idempotency_key"].value()
        form = {"customer_name": "Trần B", "phone": "+84 912 000 111", "address": "34 Đường Kiểm Tra, Quận 3",
                "payment_method": "cod", "idempotency_key": str(key),
                # Giá trị giả từ trình duyệt phải bị bỏ qua
                "total": "1", "price": "1"}
        r1 = shop.post("/thanh-toan/", form)
        self.assertEqual(r1.status_code, 302)
        order = Order.objects.get()
        self.assertEqual(order.subtotal, 318_000)
        self.assertEqual(order.shipping_fee, 30_000)
        self.assertEqual(order.total, 348_000)
        self.assertEqual(order.phone, "0912000111")
        # Bấm lại (gửi lặp) không tạo đơn mới, không trừ kho lần hai
        shop.post("/gio-hang/them/", {"variant_id": v.pk, "quantity": 1})
        shop.post("/thanh-toan/", form)
        self.assertEqual(Order.objects.count(), 1)
        self.assertEqual(Variant.objects.get(sku="SKT-01").stock, 2)
        r = shop.get(r1["Location"])
        self.assertContains(r, order.code)

        # Người khác không xem được đơn nếu không có mã + số điện thoại đúng
        other = self.client_class()
        self.assertEqual(other.get(f"/don-hang/{order.code}/").status_code, 404)
        r = other.post("/don-hang/tra-cuu/", {"code": order.code, "phone": "0999999999"})
        self.assertContains(r, "Không tìm thấy")
        r = other.post("/don-hang/tra-cuu/", {"code": order.code.lower(), "phone": "0912 000 111"})
        self.assertEqual(r.status_code, 302)
        self.assertContains(other.get(r["Location"]), "Trần B")

        # Admin xử lý đơn và lịch sử được ghi
        self.client.post(f"/admin/don-hang/{order.pk}/trang-thai/", {"status": "confirmed"})
        self.client.post(f"/admin/don-hang/{order.pk}/van-chuyen/", {"carrier": "GHN", "tracking_code": "ABC123"})
        self.client.post(f"/admin/don-hang/{order.pk}/ghi-chu/", {"note": "Gói kỹ"})
        r = self.client.get(f"/admin/don-hang/{order.pk}/")
        self.assertContains(r, "ABC123")
        self.assertContains(r, "Gói kỹ")
        self.assertEqual(order.events.count(), 4)

        # Sửa giá sau khi đặt không đổi đơn cũ, có lịch sử giá
        data = {"name": p.name, "slug": p.slug, "code": p.code, "status": "published", "price_reason": "Tăng giá NCC", **formset_data([
            {"id": v.pk, "value": "01 Đỏ", "sku": "SKT-01", "price": "179000", "compare_at_price": "", "stock": "2",
             "orig_stock": "2", "orig_price": "159000", "orig_compare": "199000", "is_active": "on", "sort_order": "0"},
            {"id": Variant.objects.get(sku="SKT-02").pk, "value": "02 Cam", "sku": "SKT-02", "price": "169000",
             "compare_at_price": "", "stock": "0", "orig_stock": "0", "orig_price": "169000", "orig_compare": "",
             "is_active": "on", "sort_order": "1"},
        ], initial=2)}
        r = self.client.post(f"/admin/san-pham/{p.pk}/", data)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Variant.objects.get(sku="SKT-01").price, 179_000)
        h = PriceHistory.objects.filter(variant__sku="SKT-01", source="manual").get()
        self.assertEqual((h.old_price, h.new_price, h.reason, h.changed_by), (159_000, 179_000, "Tăng giá NCC", self.owner))
        item = Order.objects.get().items.get()
        self.assertEqual((item.unit_price, item.line_total), (159_000, 318_000))

        # Không xóa được sản phẩm đã có trong đơn
        self.client.post(f"/admin/san-pham/{p.pk}/xoa/")
        self.assertTrue(Product.objects.filter(pk=p.pk).exists())

    def test_stock_edit_requires_reason_and_detects_concurrent_change(self):
        p, (v,) = make_product(stock=10)
        base = {"name": p.name, "slug": p.slug, "code": p.code, "status": "published"}
        row = {"id": v.pk, "value": "", "sku": v.sku, "price": str(v.price), "compare_at_price": "", "stock": "15",
               "orig_stock": "10", "orig_price": str(v.price), "orig_compare": "", "is_active": "on", "sort_order": "0"}
        r = self.client.post(f"/admin/san-pham/{p.pk}/", {**base, **formset_data([row], initial=1)})
        self.assertContains(r, "Lý do điều chỉnh tồn kho")
        self.assertEqual(Variant.objects.get(pk=v.pk).stock, 10)
        # Có đơn chen vào giữa lúc mở form và lúc lưu
        place_order({v.pk: 1}, customer(), "cod", "", uuid.uuid4())
        r = self.client.post(f"/admin/san-pham/{p.pk}/", {**base, "stock_reason": "Nhập hàng", **formset_data([row], initial=1)})
        self.assertEqual(Variant.objects.get(pk=v.pk).stock, 9)
        r = self.client.post(f"/admin/ton-kho/{v.pk}/", {"mode": "add", "quantity": 5, "reason": "Nhập hàng", "expected": 9})
        self.assertEqual(Variant.objects.get(pk=v.pk).stock, 14)
        mv = StockMovement.objects.filter(variant=v, kind="adjust").get()
        self.assertEqual((mv.delta, mv.stock_after, mv.user, mv.reason), (5, 14, self.owner, "Nhập hàng"))

    def test_rejects_fake_image(self):
        p, _ = make_product()
        fake = SimpleUploadedFile("x.jpg", b"<?php echo 1; ?>", content_type="image/jpeg")
        self.client.post(f"/admin/san-pham/{p.pk}/anh/", {"action": "upload", "images": [fake]})
        self.assertEqual(ProductImage.objects.filter(product=p).count(), 0)
        svg = SimpleUploadedFile("x.svg", b"<svg onload='alert(1)'></svg>", content_type="image/svg+xml")
        self.client.post(f"/admin/san-pham/{p.pk}/anh/", {"action": "upload", "images": [svg]})
        self.assertEqual(ProductImage.objects.filter(product=p).count(), 0)

    def test_bulk_price_preview_then_confirm(self):
        p, (v1, v2) = make_product(variants=[("30ml", 100_000, None, 1), ("50ml", 200_000, None, 1)])
        r = self.client.post("/admin/gia/hang-loat/", {"step": "preview", "q": p.name, "operation": "sale_pct", "amount": 10,
                                                       "rounding": "1000", "reason": "Sale cuối tuần"})
        self.assertContains(r, "Xác nhận cập nhật giá")
        self.assertEqual(Variant.objects.get(pk=v1.pk).price, 100_000)  # chưa đổi khi mới xem trước
        self.client.post("/admin/gia/hang-loat/", {"step": "confirm"})
        v1.refresh_from_db()
        v2.refresh_from_db()
        self.assertEqual((v1.price, v1.compare_at_price, v2.price, v2.compare_at_price), (90_000, 100_000, 180_000, 200_000))
        self.assertEqual(PriceHistory.objects.filter(source="bulk").count(), 2)

    def _xlsx(self, rows):
        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS)
        for r in rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        return SimpleUploadedFile("sp.xlsx", buf.getvalue())

    def test_excel_import_reports_row_errors_then_imports(self):
        bad = self._xlsx([
            ["SP1", "Kem dưỡng A", "Hãng Lạ", "", "Dung tích", "30ml", "KA-30", "abc", "", 5, "Đang bán"],
            ["SP1", "Kem dưỡng A", "", "", "Dung tích", "50ml", "KA-30", 200000, 150000, 2, ""],
        ])
        r = self.client.post("/admin/excel/nhap/", {"file": bad}, follow=True)
        self.assertContains(r, "Giá bán")
        self.assertContains(r, "Trùng với dòng 2")
        self.assertContains(r, "Hãng Lạ")
        self.assertEqual(Product.objects.filter(code="SP1").count(), 0)

        good = self._xlsx([
            ["SP1", "Kem dưỡng A", "Hãng Mới", "Chăm sóc da > Kem dưỡng", "Dung tích", "30ml", "KA-30", 150000, 180000, 5, "Đang bán"],
            ["SP1", "", "", "", "", "50ml", "KA-50", 220000, "", 0, ""],
        ])
        r = self.client.post("/admin/excel/nhap/", {"file": good, "auto_create": "on"})
        self.assertEqual(Product.objects.filter(code="SP1").count(), 0)  # chưa ghi khi xem trước
        r = self.client.post(r["Location"])
        p = Product.objects.get(code="SP1")
        self.assertEqual((p.status, p.brand.name, p.category.parent.name, p.variants.count()), ("published", "Hãng Mới", "Chăm sóc da", 2))
        self.assertEqual(Variant.objects.get(sku="KA-30").stock, 5)

        # Nhập lại file sửa giá + tồn: ghi lịch sử
        again = self._xlsx([["SP1", "Kem dưỡng A", "", "", "", "30ml", "KA-30", 160000, "", 8, ""]])
        r = self.client.post("/admin/excel/nhap/", {"file": again})
        self.client.post(r["Location"])
        v = Variant.objects.get(sku="KA-30")
        self.assertEqual((v.price, v.compare_at_price, v.stock), (160_000, None, 8))
        self.assertTrue(PriceHistory.objects.filter(variant=v, source="import", old_price=150_000).exists())
        self.assertTrue(StockMovement.objects.filter(variant=v, kind="import", delta=3).exists())

    def test_excel_export_and_template(self):
        make_product(name="=CMD|' /C calc'!A0")
        r = self.client.get("/admin/excel/xuat/")
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[1]][:3], HEADERS[:3])
        self.assertEqual(ws.cell(2, 2).data_type, "s")  # không bị hiểu là công thức
        self.assertEqual(self.client.get("/admin/excel/mau/").status_code, 200)

    def test_settings_validation(self):
        r = self.client.post("/admin/cai-dat/", {"shop_name": "Kenvin Shop", "primary_color": "#B23A62", "bank_transfer_enabled": "on",
                                                 "shipping_flat_fee": 0, "low_stock_threshold": 5, "max_qty_per_item": 20, "bank_transfer_hold_hours": 0})
        self.assertContains(r, "Cần nhập đủ ngân hàng")
        self.assertFalse(SiteSettings.load().bank_transfer_enabled)
