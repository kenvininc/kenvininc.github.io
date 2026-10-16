import io
from datetime import datetime, time, timedelta

from django.contrib import messages
from django.db.models import Q
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.core.roles import has_perm, require_perm
from apps.core.utils import audit
from apps.orders.models import Order
from apps.orders.services import (
    PAYMENT_TRANSITIONS,
    TRANSITIONS,
    OrderError,
    add_note,
    change_payment,
    change_status,
    restock_returned,
    update_shipping,
)

from ..forms import OrderFilterForm, ShippingForm
from .common import paginate, xlsx_response


def _filtered(request):
    form = OrderFilterForm(request.GET or None)
    qs = Order.objects.all()
    if form.is_valid():
        cd = form.cleaned_data
        q = cd["q"].strip()
        if q:
            qs = qs.filter(Q(code__icontains=q) | Q(customer_name__icontains=q) | Q(phone__icontains=q.replace(" ", "")))
        if cd["status"]:
            qs = qs.filter(status=cd["status"])
        if cd["payment_status"]:
            qs = qs.filter(payment_status=cd["payment_status"])
        tz = timezone.get_current_timezone()
        if cd["date_from"]:
            qs = qs.filter(created_at__gte=datetime.combine(cd["date_from"], time.min, tzinfo=tz))
        if cd["date_to"]:
            qs = qs.filter(created_at__lt=datetime.combine(cd["date_to"] + timedelta(days=1), time.min, tzinfo=tz))
    return form, qs


@require_perm("view_orders")
def order_list(request):
    form, qs = _filtered(request)
    counts = {s: Order.objects.filter(status=s).count() for s in (Order.Status.NEW, Order.Status.CONFIRMED, Order.Status.PACKING)}
    return render(request, "panel/orders/list.html", {"form": form, "page_obj": paginate(request, qs), "counts": counts})


@require_perm("view_orders")
def order_detail(request, pk):
    order = get_object_or_404(Order.objects.prefetch_related("items"), pk=pk)
    events = order.events.select_related("user")
    status_labels = dict(Order.Status.choices)
    pay_labels = dict(Order.PaymentStatus.choices)
    return render(
        request,
        "panel/orders/detail.html",
        {
            "order": order,
            "events": events,
            "next_statuses": [(s, status_labels[s]) for s in TRANSITIONS[order.status]],
            "next_payments": [(s, pay_labels[s]) for s in PAYMENT_TRANSITIONS[order.payment_status]],
            "shipping_form": ShippingForm(initial={"carrier": order.carrier, "tracking_code": order.tracking_code}),
            "status_labels": status_labels,
            "pay_labels": pay_labels,
            "can_manage": has_perm(request.user, "manage_orders"),
            "can_pay": has_perm(request.user, "manage_payments"),
        },
    )


@require_POST
@require_perm("manage_orders")
def order_status(request, pk):
    new = request.POST.get("status", "")
    note = (request.POST.get("note") or "").strip()[:1000]
    if new not in Order.Status.values:
        messages.error(request, "Trạng thái không hợp lệ.")
        return redirect("panel:order_detail", pk=pk)
    if new == Order.Status.CANCELLED and not note:
        messages.error(request, "Vui lòng ghi lý do hủy đơn.")
        return redirect("panel:order_detail", pk=pk)
    try:
        order = change_status(pk, new, request.user, note=note, restock_return=request.POST.get("restock") == "1")
    except OrderError as e:
        messages.error(request, str(e))
    else:
        audit(request, "order_status", order, f"{order.code} → {order.get_status_display()}")
        messages.success(request, f"Đơn {order.code}: {order.get_status_display()}.")
        if new == Order.Status.CANCELLED and order.payment_status == Order.PaymentStatus.PAID:
            messages.warning(request, "Đơn đã thanh toán: nhớ hoàn tiền cho khách và cập nhật trạng thái thanh toán.")
    return redirect("panel:order_detail", pk=pk)


@require_POST
@require_perm("manage_orders")
def order_restock(request, pk):
    try:
        order = restock_returned(pk, request.user)
    except OrderError as e:
        messages.error(request, str(e))
    else:
        audit(request, "order_restock", order, order.code)
        messages.success(request, "Đã nhập lại kho hàng hoàn.")
    return redirect("panel:order_detail", pk=pk)


@require_POST
@require_perm("manage_payments")
def order_payment(request, pk):
    new = request.POST.get("payment_status", "")
    if new not in Order.PaymentStatus.values:
        messages.error(request, "Trạng thái thanh toán không hợp lệ.")
        return redirect("panel:order_detail", pk=pk)
    try:
        order = change_payment(pk, new, request.user, note=(request.POST.get("note") or "")[:1000])
    except OrderError as e:
        messages.error(request, str(e))
    else:
        audit(request, "order_payment", order, f"{order.code} → {order.get_payment_status_display()}")
        messages.success(request, f"Đơn {order.code}: {order.get_payment_status_display()}.")
    return redirect("panel:order_detail", pk=pk)


@require_POST
@require_perm("manage_orders")
def order_shipping(request, pk):
    form = ShippingForm(request.POST)
    if form.is_valid():
        order = update_shipping(pk, form.cleaned_data["carrier"].strip(), form.cleaned_data["tracking_code"].strip(), request.user)
        audit(request, "order_shipping", order, f"{order.code}: {order.carrier} {order.tracking_code}")
        messages.success(request, "Đã lưu thông tin vận chuyển.")
    else:
        messages.error(request, "Thông tin vận chuyển không hợp lệ.")
    return redirect("panel:order_detail", pk=pk)


@require_POST
@require_perm("manage_orders")
def order_note(request, pk):
    get_object_or_404(Order, pk=pk)
    try:
        add_note(pk, (request.POST.get("note") or "")[:2000], request.user)
    except OrderError as e:
        messages.error(request, str(e))
    else:
        messages.success(request, "Đã thêm ghi chú nội bộ.")
    return redirect("panel:order_detail", pk=pk)


@require_perm("view_orders")
def order_print(request, pk):
    order = get_object_or_404(Order.objects.prefetch_related("items"), pk=pk)
    kind = request.GET.get("loai", "dong-goi")
    return render(request, "panel/orders/print.html", {"orders": [order], "kind": kind})


@require_perm("view_orders")
def order_print_many(request):
    ids = [int(i) for i in request.GET.getlist("ids") if i.isdigit()][:100]
    orders = Order.objects.filter(pk__in=ids).prefetch_related("items")
    return render(request, "panel/orders/print.html", {"orders": orders, "kind": request.GET.get("loai", "dong-goi")})


def _cell(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        c.data_type = "s"


@require_perm("view_orders")
def order_export(request):
    _form, qs = _filtered(request)
    qs = qs.prefetch_related("items")[:20000]
    wb = Workbook()
    ws = wb.active
    ws.title = "Don hang"
    head = ["Mã đơn", "Ngày đặt", "Khách hàng", "Điện thoại", "Email", "Địa chỉ", "Ghi chú khách", "Trạng thái",
            "Thanh toán", "Phương thức", "Tiền hàng", "Giảm giá", "Mã giảm giá", "Phí giao hàng", "Tổng thanh toán",
            "ĐV vận chuyển", "Mã vận đơn", "Ngày hoàn tất"]
    ws.append(head)
    ws2 = wb.create_sheet("Chi tiet")
    ws2.append(["Mã đơn", "Sản phẩm", "Phân loại", "SKU", "Đơn giá", "Số lượng", "Thành tiền"])
    for c in list(ws[1]) + list(ws2[1]):
        c.font = Font(bold=True)
    r, r2 = 2, 2
    for o in qs:
        values = [o.code, timezone.localtime(o.created_at).strftime("%d/%m/%Y %H:%M"), o.customer_name, o.phone, o.email,
                  o.address, o.customer_note, o.get_status_display(), o.get_payment_status_display(),
                  o.get_payment_method_display(), o.subtotal, o.discount_total, o.coupon_code, o.shipping_fee, o.total,
                  o.carrier, o.tracking_code,
                  timezone.localtime(o.completed_at).strftime("%d/%m/%Y %H:%M") if o.completed_at else ""]
        for i, v in enumerate(values, 1):
            _cell(ws, r, i, v)
        r += 1
        for it in o.items.all():
            for i, v in enumerate([o.code, it.product_name, it.variant_label, it.sku, it.unit_price, it.quantity, it.line_total], 1):
                _cell(ws2, r2, i, v)
            r2 += 1
    buf = io.BytesIO()
    wb.save(buf)
    audit(request, "excel_export_orders", summary=f"{r - 2} đơn")
    return xlsx_response(buf.getvalue(), f"don-hang-{timezone.localtime():%Y%m%d-%H%M}.xlsx")
