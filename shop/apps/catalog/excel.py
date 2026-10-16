"""Nhập/xuất sản phẩm bằng Excel.

Mỗi dòng là 1 biến thể (SKU). Các dòng cùng "Mã sản phẩm" thuộc cùng 1 sản phẩm.
Nhập gồm 2 bước: kiểm tra (không ghi gì) → xem trước → xác nhận mới ghi, trong 1 transaction.
"""

import io

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Attribute, Brand, Category, PriceHistory, Product, StockMovement, Variant, unique_slug
from .services import ConcurrentChange, adjust_stock, set_price

COLUMNS = [
    ("code", "Mã sản phẩm", True, "Mã nội bộ. Các dòng cùng mã = các biến thể của 1 sản phẩm."),
    ("name", "Tên sản phẩm", True, "Bắt buộc ở dòng đầu tiên của mỗi sản phẩm."),
    ("brand", "Thương hiệu", False, "Tên thương hiệu đã có trong hệ thống."),
    ("category", "Danh mục", False, "Tên danh mục. Danh mục con ghi dạng: Cha > Con"),
    ("attribute", "Loại biến thể", False, "Ví dụ: Dung tích, Màu sắc. Để trống nếu chỉ có 1 loại."),
    ("value", "Giá trị biến thể", False, "Ví dụ: 50ml, 01 Đỏ cam."),
    ("sku", "SKU", True, "Mã biến thể, không trùng."),
    ("price", "Giá bán", True, "Số nguyên, đơn vị đồng. Ví dụ 159000"),
    ("compare_at_price", "Giá gốc", False, "Để trống nếu không giảm giá. Phải lớn hơn giá bán."),
    ("stock", "Tồn kho", True, "Số lượng tồn thực tế (số nguyên ≥ 0)."),
    ("status", "Trạng thái", False, "Nháp / Đang bán / Ẩn. Để trống = Nháp (sản phẩm mới) hoặc giữ nguyên."),
    ("short_description", "Mô tả ngắn", False, ""),
    ("description", "Mô tả", False, ""),
    ("ingredients", "Thành phần", False, ""),
    ("benefits", "Công dụng", False, ""),
    ("usage", "Hướng dẫn sử dụng", False, ""),
]
HEADERS = [c[1] for c in COLUMNS]
KEYS = [c[0] for c in COLUMNS]
STATUS_IN = {"nháp": "draft", "nhap": "draft", "đang bán": "published", "dang ban": "published", "ẩn": "hidden", "an": "hidden"}
STATUS_OUT = {"draft": "Nháp", "published": "Đang bán", "hidden": "Ẩn"}
PRODUCT_FIELDS = ["name", "brand", "category", "attribute", "status", "short_description", "description", "ingredients", "benefits", "usage"]
MAX_ROWS = 5000


def _safe(ws, row, col, value):
    """Ghi ô, không để chuỗi bắt đầu bằng '=' bị Excel hiểu là công thức."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        cell.data_type = "s"
    return cell


def _header(ws, headers, widths=None):
    fill = PatternFill("solid", fgColor="F3DCE4")
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 18)
    ws.freeze_panes = "A2"


def _guide_sheet(wb):
    ws = wb.create_sheet("Hướng dẫn")
    ws.append(["Cột", "Bắt buộc", "Ghi chú"])
    for _key, label, req, note in COLUMNS:
        ws.append([label, "Có" if req else "", note])
    ws.append([])
    ws.append(["Quy tắc chung"])
    for line in [
        "Mỗi dòng là 1 biến thể (SKU). Sản phẩm 1 loại thì chỉ có 1 dòng.",
        "Mã sản phẩm đã có trong hệ thống → cập nhật sản phẩm đó. Chưa có → tạo mới.",
        "Khi cập nhật: ô để trống = giữ nguyên giá trị cũ (trừ Giá gốc: để trống = bỏ giá gốc).",
        "Tồn kho là số lượng thực tế sau khi nhập, hệ thống tự tính chênh lệch và ghi lịch sử.",
        "Ảnh sản phẩm không nhập qua Excel, upload trong trang sửa sản phẩm.",
        "Có lỗi ở bất kỳ dòng nào thì không dòng nào được ghi. Sửa file rồi tải lên lại.",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 80
    for c in ws[1]:
        c.font = Font(bold=True)


def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "San pham"
    _header(ws, HEADERS, {1: 16, 2: 36, 7: 18, 13: 40})
    _guide_sheet(wb)
    return _to_bytes(wb)


def export_products(products):
    wb = Workbook()
    ws = wb.active
    ws.title = "San pham"
    _header(ws, HEADERS, {1: 16, 2: 36, 7: 18, 13: 40})
    r = 2
    for p in products.select_related("brand", "category", "category__parent", "variant_attribute").prefetch_related("variants"):
        cat = ""
        if p.category_id:
            cat = f"{p.category.parent.name} > {p.category.name}" if p.category.parent_id else p.category.name
        for v in p.variants.all():
            values = {
                "code": p.code, "name": p.name, "brand": p.brand.name if p.brand_id else "", "category": cat,
                "attribute": p.variant_attribute.name if p.variant_attribute_id else "", "value": v.value,
                "sku": v.sku, "price": v.price, "compare_at_price": v.compare_at_price, "stock": v.stock,
                "status": STATUS_OUT[p.status], "short_description": p.short_description, "description": p.description,
                "ingredients": p.ingredients, "benefits": p.benefits, "usage": p.usage,
            }
            for i, k in enumerate(KEYS, 1):
                _safe(ws, r, i, values[k])
            r += 1
    _guide_sheet(wb)
    return _to_bytes(wb)


def _to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- Đọc file


def _str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _money(v):
    s = _str(v).replace(".", "").replace(",", "").replace(" ", "").replace("đ", "").replace("₫", "")
    if s == "":
        return None
    if not s.isdigit():
        raise ValueError
    return int(s)


def _find_category(label, cats):
    label = label.strip()
    if ">" in label:
        parent, child = [x.strip() for x in label.split(">", 1)]
        return cats.get((parent.lower(), child.lower()))
    return cats.get((None, label.lower()))


def parse_import(fileobj, auto_create=False):
    """Đọc và kiểm tra file. Trả về (plan, summary, errors). Không ghi database."""
    errors = []

    def err(row, col, msg):
        errors.append({"row": row, "col": col, "msg": msg})

    try:
        wb = load_workbook(fileobj, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - mọi lỗi đọc file đều báo chung
        return None, None, [{"row": 0, "col": "", "msg": "Không đọc được file. Hãy dùng file .xlsx theo mẫu."}]
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [(_str(h)) for h in next(rows, [])]
    missing = [h for h in HEADERS if h not in header]
    if missing:
        return None, None, [{"row": 1, "col": "", "msg": "Thiếu cột: " + ", ".join(missing) + ". Hãy dùng file mẫu."}]
    idx = {KEYS[HEADERS.index(h)]: header.index(h) for h in HEADERS}

    brands = {b.name.lower(): b.pk for b in Brand.objects.all()}
    cats = {}
    for c in Category.objects.select_related("parent"):
        cats[(c.parent.name.lower() if c.parent_id else None, c.name.lower())] = c.pk
    attrs = {a.name.lower(): a.pk for a in Attribute.objects.all()}
    existing_products = {p.code: p for p in Product.objects.all()}
    existing_variants = {v.sku: v for v in Variant.objects.select_related("product")}

    groups = {}
    seen_sku = {}
    count = 0
    for rnum, raw in enumerate(rows, start=2):
        if raw is None or all(_str(x) == "" for x in raw):
            continue
        count += 1
        if count > MAX_ROWS:
            err(rnum, "", f"File vượt quá {MAX_ROWS} dòng, hãy chia nhỏ file.")
            break
        cell = {k: (raw[i] if i < len(raw) else None) for k, i in idx.items()}
        code = _str(cell["code"])
        sku = _str(cell["sku"])
        if not code:
            err(rnum, "Mã sản phẩm", "Bắt buộc.")
            continue
        if len(code) > 60:
            err(rnum, "Mã sản phẩm", "Tối đa 60 ký tự.")
        if not sku:
            err(rnum, "SKU", "Bắt buộc.")
        elif sku in seen_sku:
            err(rnum, "SKU", f"Trùng với dòng {seen_sku[sku]}.")
        elif len(sku) > 64:
            err(rnum, "SKU", "Tối đa 64 ký tự.")
        else:
            seen_sku[sku] = rnum
            ev = existing_variants.get(sku)
            if ev and ev.product.code != code:
                err(rnum, "SKU", f"SKU đã thuộc sản phẩm khác (mã {ev.product.code}).")
        try:
            price = _money(cell["price"])
            if price is None:
                err(rnum, "Giá bán", "Bắt buộc.")
            elif price <= 0:
                err(rnum, "Giá bán", "Phải lớn hơn 0.")
        except ValueError:
            price = None
            err(rnum, "Giá bán", "Phải là số nguyên (đồng).")
        try:
            compare = _money(cell["compare_at_price"])
            if compare is not None and price is not None and compare <= price:
                err(rnum, "Giá gốc", "Phải lớn hơn giá bán (hoặc để trống).")
        except ValueError:
            compare = None
            err(rnum, "Giá gốc", "Phải là số nguyên (đồng).")
        try:
            stock = _money(cell["stock"])
            if stock is None:
                err(rnum, "Tồn kho", "Bắt buộc.")
        except ValueError:
            stock = None
            err(rnum, "Tồn kho", "Phải là số nguyên ≥ 0.")
        status_raw = _str(cell["status"]).lower()
        status = STATUS_IN.get(status_raw) if status_raw else ""
        if status_raw and status is None:
            err(rnum, "Trạng thái", "Chỉ nhận: Nháp, Đang bán, Ẩn.")

        g = groups.setdefault(code, {"rows": [], "fields": {}, "first_row": rnum})
        fields = {
            "name": _str(cell["name"]), "brand": _str(cell["brand"]), "category": _str(cell["category"]),
            "attribute": _str(cell["attribute"]), "status": status or "",
            "short_description": _str(cell["short_description"]), "description": _str(cell["description"]),
            "ingredients": _str(cell["ingredients"]), "benefits": _str(cell["benefits"]), "usage": _str(cell["usage"]),
        }
        for k, v in fields.items():
            if not v:
                continue
            if k in g["fields"] and g["fields"][k] != v:
                err(rnum, dict(zip(KEYS, HEADERS))[k], f"Khác với dòng {g['first_row']} của cùng mã sản phẩm.")
            g["fields"].setdefault(k, v)
        g["rows"].append(
            {"row": rnum, "sku": sku, "value": _str(cell["value"])[:80], "price": price, "compare": compare, "stock": stock}
        )

    if count == 0 and not errors:
        err(0, "", "File không có dòng dữ liệu nào.")

    plan = {"products": [], "create_brands": [], "create_categories": [], "create_attributes": []}
    summary = {"products_new": 0, "products_updated": 0, "variants_new": 0, "variants_updated": 0, "price_changes": 0, "stock_changes": 0}
    for code, g in groups.items():
        f = g["fields"]
        row = g["first_row"]
        prod = existing_products.get(code)
        if not prod and not f.get("name"):
            err(row, "Tên sản phẩm", "Bắt buộc với sản phẩm mới.")
        if len(f.get("name", "")) > 255:
            err(row, "Tên sản phẩm", "Tối đa 255 ký tự.")
        if len(f.get("short_description", "")) > 300:
            err(row, "Mô tả ngắn", "Tối đa 300 ký tự.")
        for key, table, label, bucket in (
            ("brand", brands, "Thương hiệu", "create_brands"),
            ("attribute", attrs, "Loại biến thể", "create_attributes"),
        ):
            name = f.get(key)
            if name and name.lower() not in table:
                if auto_create:
                    if name not in plan[bucket]:
                        plan[bucket].append(name)
                else:
                    err(row, label, f"“{name}” chưa có trong hệ thống.")
        if f.get("category") and _find_category(f["category"], cats) is None:
            if auto_create:
                if f["category"] not in plan["create_categories"]:
                    plan["create_categories"].append(f["category"])
            else:
                err(row, "Danh mục", f"“{f['category']}” chưa có trong hệ thống.")
        variants_existing_count = prod.variants.count() if prod else 0
        multi = len(g["rows"]) > 1 or variants_existing_count > 1
        values = [r["value"] for r in g["rows"]]
        if multi and any(not v for v in values):
            err(row, "Giá trị biến thể", "Sản phẩm có nhiều biến thể thì mỗi dòng phải có giá trị biến thể.")
        if len(set(values)) != len(values):
            err(row, "Giá trị biến thể", "Các biến thể của cùng sản phẩm bị trùng giá trị.")

        if prod:
            summary["products_updated"] += 1
        else:
            summary["products_new"] += 1
        for r in g["rows"]:
            ev = existing_variants.get(r["sku"])
            if ev:
                summary["variants_updated"] += 1
                r["variant_id"] = ev.pk
                r["expected_stock"] = ev.stock
                r["expected_price"] = [ev.price, ev.compare_at_price]
                if r["price"] is not None and (ev.price, ev.compare_at_price) != (r["price"], r["compare"]):
                    summary["price_changes"] += 1
                if r["stock"] is not None and ev.stock != r["stock"]:
                    summary["stock_changes"] += 1
            else:
                summary["variants_new"] += 1
        plan["products"].append({"code": code, "product_id": prod.pk if prod else None, "fields": f, "rows": g["rows"]})
    errors.sort(key=lambda e: e["row"])
    return plan, summary, errors


def _resolve_category(label, user):
    if ">" in label:
        parent_name, child_name = [x.strip() for x in label.split(">", 1)]
        parent = Category.objects.filter(parent__isnull=True, name__iexact=parent_name).first()
        if not parent:
            parent = Category.objects.create(name=parent_name, slug=unique_slug(Category, parent_name))
        child = Category.objects.filter(parent=parent, name__iexact=child_name).first()
        if not child:
            child = Category.objects.create(name=child_name, parent=parent, slug=unique_slug(Category, child_name))
        return child
    cat = Category.objects.filter(parent__isnull=True, name__iexact=label).first() or Category.objects.filter(
        name__iexact=label
    ).first()
    if not cat:
        cat = Category.objects.create(name=label, slug=unique_slug(Category, label))
    return cat


def apply_import(plan, user):
    """Ghi dữ liệu theo plan đã kiểm tra. Raise ConcurrentChange nếu dữ liệu đổi kể từ lúc xem trước."""
    with transaction.atomic():
        for name in plan["create_brands"]:
            if not Brand.objects.filter(name__iexact=name).exists():
                Brand.objects.create(name=name, slug=unique_slug(Brand, name))
        for name in plan["create_attributes"]:
            if not Attribute.objects.filter(name__iexact=name).exists():
                Attribute.objects.create(name=name)
        for item in plan["products"]:
            f = item["fields"]
            if item["product_id"]:
                product = Product.objects.select_for_update().get(pk=item["product_id"])
            else:
                if Product.objects.filter(code=item["code"]).exists():
                    raise ConcurrentChange(f"Mã sản phẩm {item['code']} vừa được tạo bởi thao tác khác.")
                product = Product(code=item["code"], created_by=user, status=Product.Status.DRAFT)
            for key in ("name", "short_description", "description", "ingredients", "benefits", "usage"):
                if f.get(key):
                    setattr(product, key, f[key])
            if f.get("status"):
                product.status = f["status"]
            if f.get("brand"):
                product.brand = Brand.objects.get(name__iexact=f["brand"])
            if f.get("category"):
                product.category = _resolve_category(f["category"], user)
            if f.get("attribute"):
                product.variant_attribute = Attribute.objects.get(name__iexact=f["attribute"])
            if not product.slug:
                product.slug = unique_slug(Product, product.name)
            product.save()

            for r in item["rows"]:
                if r.get("variant_id"):
                    v = Variant.objects.select_for_update().get(pk=r["variant_id"])
                    if v.product_id != product.pk:
                        raise ConcurrentChange(f"SKU {v.sku} đã bị chuyển sang sản phẩm khác.")
                    if r["value"] and r["value"] != v.value:
                        v.value = r["value"]
                        v.save(update_fields=["value"])
                    set_price(
                        v, r["price"], r["compare"], user, PriceHistory.Source.IMPORT,
                        reason="Nhập Excel", expected=r["expected_price"],
                    )
                    adjust_stock(
                        v.pk, r["stock"] - r["expected_stock"], StockMovement.Kind.IMPORT, user=user,
                        reason="Nhập Excel", expected=r["expected_stock"],
                    )
                else:
                    if Variant.objects.filter(sku=r["sku"]).exists():
                        raise ConcurrentChange(f"SKU {r['sku']} vừa được tạo bởi thao tác khác.")
                    v = Variant.objects.create(
                        product=product, sku=r["sku"], value=r["value"], price=r["price"],
                        compare_at_price=r["compare"], stock=0,
                    )
                    PriceHistory.objects.create(
                        variant=v, old_price=None, new_price=v.price, new_compare_at_price=v.compare_at_price,
                        source=PriceHistory.Source.IMPORT, reason="Tạo từ Excel", changed_by=user,
                    )
                    if r["stock"]:
                        adjust_stock(v.pk, r["stock"], StockMovement.Kind.IMPORT, user=user, reason="Nhập Excel")
            product.refresh_search_text()
    return timezone.now()
